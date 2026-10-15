import sys, os
from pptx import Presentation
from pptx.util import Emu

src = r"D:\Documents\Prof_Curriculur\DIP\Chap5 图像复原与重建-final.pptx"
out = r"D:\Note\OB_Note\DIP\_tmp_ch5\ch5_raw.txt"
os.makedirs(os.path.dirname(out), exist_ok=True)

prs = Presentation(src)
lines = []
lines.append(f"TOTAL SLIDES: {len(prs.slides)}")
lines.append(f"SLIDE SIZE: {prs.slide_width} x {prs.slide_height}")

def walk(shapes, indent="", depth=0):
    res = []
    for sh in shapes:
        try:
            st = sh.shape_type
        except Exception:
            st = None
        nm = sh.name
        if sh.shape_type is not None and str(sh.shape_type).startswith("GROUP"):
            res.append(f"{indent}[GROUP {nm}]")
            res.extend(walk(sh.shapes, indent + "  ", depth + 1))
            continue
        if sh.has_text_frame:
            txt = sh.text_frame.text
            if txt and txt.strip():
                res.append(f"{indent}[{nm}] {txt}")
            continue
        if getattr(sh, "has_table", False) and sh.has_table:
            res.append(f"{indent}[TABLE {nm}]")
            for r in sh.table.rows:
                cells = [c.text.replace("\n", " ") for c in r.cells]
                res.append(f"{indent}  | " + " | ".join(cells) + " |")
            continue
        if sh.shape_type is not None and "PICTURE" in str(sh.shape_type):
            res.append(f"{indent}[PICTURE {nm}]")
            continue
        if sh.shape_type is not None and "CHART" in str(sh.shape_type):
            res.append(f"{indent}[CHART {nm}]")
            continue
        res.append(f"{indent}[SHAPE {nm} type={st}]")
    return res

for i, slide in enumerate(prs.slides, 1):
    lines.append("")
    lines.append(f"===== SLIDE {i} =====")
    try:
        if slide.has_notes_slide:
            nt = slide.notes_slide.notes_text_frame.text
            if nt and nt.strip():
                lines.append(f"[NOTES] {nt}")
    except Exception:
        pass
    try:
        lines.extend(walk(slide.shapes))
    except Exception as e:
        lines.append(f"[ERROR walking shapes: {e}]")

with open(out, "w", encoding="utf-8") as f:
    f.write("\n".join(lines))
print("OK slides:", len(prs.slides), "chars:", sum(len(x) for x in lines))
print("written:", out)
